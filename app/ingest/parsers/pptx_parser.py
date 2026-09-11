from __future__ import annotations

from pathlib import Path

from app.ingest.classify import apply_structure, content_type_for
from app.ingest.models import Block, ParsedDocument


def parse_pptx(path: Path, *, use_ocr: bool = True) -> ParsedDocument:
    del use_ocr
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError("python-pptx is not installed.") from exc

    pres = Presentation(str(path))
    blocks: list[Block] = []
    native_pages = 0
    empty_pages = 0

    for index, slide in enumerate(pres.slides, start=1):
        slide_blocks: list[Block] = []
        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                lines = []
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
                if lines:
                    slide_blocks.append(
                        Block(
                            type="table",
                            text="\n".join(lines),
                            page=index,
                            content_type="table",
                        )
                    )
                continue

            if shape.has_text_frame:
                texts = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
                if not texts:
                    continue
                joined = "\n".join(texts)
                is_title = False
                try:
                    is_title = bool(shape.is_placeholder and shape.placeholder_format.idx in {0, 1})
                except Exception:
                    is_title = False
                if is_title or (index == 1 and len(joined) < 80 and len(slide_blocks) == 0):
                    slide_blocks.append(
                        Block(
                            type="heading",
                            text=joined,
                            page=index,
                            level=1 if index == 1 else 2,
                            content_type="text",
                        )
                    )
                else:
                    slide_blocks.append(
                        Block(
                            type="paragraph",
                            text=joined,
                            page=index,
                            content_type=content_type_for("paragraph"),
                        )
                    )
                continue

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                slide_blocks.append(
                    Block(
                        type="figure",
                        text=f"[figure on slide {index}]",
                        page=index,
                        content_type="figure",
                    )
                )

        if slide_blocks:
            native_pages += 1
            blocks.extend(slide_blocks)
        else:
            empty_pages += 1

    blocks = apply_structure(blocks)
    title = next((b.text for b in blocks if b.type == "heading"), path.stem)
    return ParsedDocument(
        source_path=str(path),
        source_name=path.name,
        format="pptx",
        title=title,
        page_count=len(pres.slides),
        native_pages=native_pages,
        empty_pages=empty_pages,
        blocks=blocks,
    )
