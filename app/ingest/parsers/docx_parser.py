from __future__ import annotations

from pathlib import Path

from app.ingest.classify import apply_structure, content_type_for
from app.ingest.models import Block, ParsedDocument


def parse_docx(path: Path, *, use_ocr: bool = True) -> ParsedDocument:
    del use_ocr  # DOCX is native text; images-only docs are rare here.
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed.") from exc

    document = docx.Document(str(path))
    blocks: list[Block] = []

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style else ""
        level = None
        block_type = "paragraph"
        if style.lower().startswith("heading"):
            block_type = "heading"
            digits = "".join(ch for ch in style if ch.isdigit())
            level = int(digits) if digits else 1
        elif "list" in style.lower():
            block_type = "list_item"
        blocks.append(
            Block(
                type=block_type,
                text=text,
                level=level,
                content_type=content_type_for(block_type),
            )
        )

    for table in document.tables:
        lines = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        if lines:
            blocks.append(
                Block(
                    type="table",
                    text="\n".join(lines),
                    content_type="table",
                    meta={"rows": len(lines)},
                )
            )

    blocks = apply_structure(blocks)
    title = next((b.text for b in blocks if b.type == "heading"), path.stem)
    return ParsedDocument(
        source_path=str(path),
        source_name=path.name,
        format="docx",
        title=title,
        page_count=1,
        native_pages=1 if blocks else 0,
        empty_pages=0 if blocks else 1,
        blocks=blocks,
    )
